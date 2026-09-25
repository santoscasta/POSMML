from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = '/Users/santoscasta/Documents/POSMML/output/documentacion-shopify/Mapa de datos Shopify POS MML.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.70); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
for name, font, size in [('Normal','Aptos',10.5), ('Title','Aptos Display',25), ('Heading 1','Aptos',16), ('Heading 2','Aptos',12)]:
    s=doc.styles[name]; s.font.name=font; s.font.size=Pt(size); s._element.rPr.rFonts.set(qn('w:eastAsia'),font)
    if name != 'Normal': s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
def cellstyle(c, header=False, alt=False):
    tcpr=c._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for side in ('top','start','bottom','end'):
        x=OxmlElement('w:'+side); x.set(qn('w:w'),'90'); x.set(qn('w:type'),'dxa'); m.append(x)
    tcpr.append(m); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'163A5F' if header else ('F3F7FA' if alt else 'FFFFFF')); tcpr.append(shd)
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.04
        for r in p.runs: r.font.size=Pt(8.6); r.font.color.rgb=RGBColor(255,255,255) if header else RGBColor(0,0,0); r.font.bold=header
def tbl(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,x in enumerate(headers): t.rows[0].cells[i].text=x; t.rows[0].cells[i].width=Inches(widths[i]); cellstyle(t.rows[0].cells[i],True)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=x; cells[i].width=Inches(widths[i]); cellstyle(cells[i],False,n%2==1)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
def h(x, level=1):
    p=doc.add_paragraph(x,style='Heading '+str(level)); p.paragraph_format.space_before=Pt(11 if level==1 else 7); p.paragraph_format.space_after=Pt(5)
def para(x, lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    if lead: r=p.add_run(lead); r.bold=True
    p.add_run(x)
def bullet(x):
    p=doc.add_paragraph(style='List Bullet'); p.add_run(x); p.paragraph_format.space_after=Pt(2)

p=doc.add_paragraph('Mapa de datos Shopify POS MML',style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
p=doc.add_paragraph('Inventario de recursos, campos y relaciones que utiliza la aplicación'); p.runs[0].italic=True; p.runs[0].font.size=Pt(12)
para('Este documento describe cómo está organizada la información que usa POS MML. La fuente principal es Shopify Admin GraphQL API versión 2025 04. Shopify guarda el catálogo, pedidos, clientes, inventario, vales, reembolsos y cumplimientos. La aplicación añade categorías y sesiones de caja mediante metaobjetos, y conserva un diario técnico local para completar operaciones de forma segura.')
para('Alcance revisado: consultas y mutaciones del frontend y servidor. No es una exportación de una tienda concreta ni presupone que todos los recursos o campos existentes en Shopify estén usados por la aplicación.','Cómo leerlo. ')

h('Resumen de arquitectura')
tbl(['Capa','Dónde vive','Qué contiene','Uso en la app'],[
['Catálogo','Shopify Product y ProductVariant','Artículos, variantes, precios, SKU, código de barras, imágenes y stock','TPV, búsqueda, escáner y cambios'],
['Clasificación','Metaobject categoria y metafield de producto','Árbol de categorías y pertenencia de productos','Navegación por categorías'],
['Venta','Shopify DraftOrder y Order','Borrador, pedido final, líneas, descuento, cliente, nota, tags y estado','Cobro, historial, ticket y dashboard'],
['Cobros y devolución','Order Transaction y Refund','Pago manual, reembolso, líneas devueltas y reposición','Caja, devoluciones y cambios'],
['Vales','GiftCard y GiftCardTransaction','Código, saldo, cliente, anotaciones y movimientos','Emisión, canje, consulta y cancelación'],
['Caja','Metaobject $app:pos_session','Apertura, cierre, arqueo y observaciones','Sesiones y KPIs'],
['Fiabilidad técnica','Volumen local POS_DATA_DIR','journal.json de operaciones y movimientos','Reintentos seguros y contabilidad de caja']], [1.0,1.45,2.65,1.45])
h('Relaciones principales')
para('Producto 1 a N Variante. Cada línea de pedido referencia una variante. Un pedido puede tener un cliente, muchas líneas, transacciones, reembolsos y órdenes de cumplimiento. Un vale puede pertenecer a un cliente, originarse en un pedido y tener muchas transacciones. La sesión de caja se relaciona con los movimientos mediante el diario local; Shopify no almacena session_id como metafield en las ventas nuevas de esta versión.')
tbl(['Origen','Relación','Destino','Clave utilizada'],[
['Product','1 a N','ProductVariant','variantId en líneas de borrador y pedido'],['Product','N a N lógico','Metaobject categoria','metafield custom.categorias con IDs de categoría en JSON'],['Order','N a 1 opcional','Customer','customerId'],['Order','1 a N','LineItem, Transaction, Refund y FulfillmentOrder','IDs nativos de Shopify'],['GiftCard','N a 1 opcional','Customer','customerId'],['Movimiento local','N a 1','Metaobject pos_session','sessionId con GID del metaobjeto']], [1.2,1.1,2.45,2.0])

h('Catálogo y categorías')
h('Productos y variantes',2)
tbl(['Recurso','Campo Shopify usado','Tipo o formato','Finalidad'],[
['Product','id','GID','Identificador'],['Product','title y productType','Texto','Nombre y tipo comercial'],['Product','status','Enum','Solo activos pueden entrar en un cambio'],['Product','featuredImage url altText','Imagen','Miniatura'],['Product','totalInventory','Número','Stock agregado'],['ProductVariant','id y title','GID y texto','Selección de variante'],['ProductVariant','price','Decimal como texto','Precio base y validación'],['ProductVariant','sku y barcode','Texto','Búsqueda y escáner'],['ProductVariant','inventoryQuantity','Número','Stock mostrado'],['ProductVariant','selectedOptions name value','Lista','Talla, color u opciones'],['ProductVariant','image url altText','Imagen','Imagen específica']], [1.25,2.1,1.35,2.1])
para('Lectura: productos ordenados por título y paginados; se cargan hasta 20 variantes por producto. La búsqueda por código de barras consulta productos de Shopify y verifica las variantes recibidas.')
h('Categorías propias',2)
tbl(['Elemento','Ubicación','Campos usados','Regla'],[
['Definición','Metaobject tipo categoria','id, displayName, nombre, imagen, padre','nombre tiene prioridad; si falta, displayName'],['Jerarquía','Campo padre','GID de metaobjeto padre o vacío','Debe existir y no puede formar ciclos'],['Imagen','Campo imagen','URL http o https','Se ignora si no es URL web'],['Asignación','Product metafield custom.categorias','Array JSON de IDs','La app rechaza valores que no sean una lista de textos']], [1.3,1.75,2.25,1.5])

h('Pedidos, cobros y preparación')
para('Una venta se crea primero como DraftOrder y se completa con paymentPending true. Solo después de validar los vales se marca como pagada. El precio de cada línea se envía como priceOverride en euros; el servidor detiene el cobro si el total calculado por Shopify difiere del importe a cobrar.')
tbl(['Recurso','Campos leídos o escritos','Uso'],[
['DraftOrder','lineItems variantId quantity priceOverride; customerId; note; appliedDiscount; tags','Crea borrador. Tags: POS MML y pos-op para recuperación.'],['Order','id, name, createdAt, cancelledAt, note, tags','Historial y trazabilidad'],['Importes','totalPriceSet, subtotalPriceSet, totalTaxSet, totalDiscountsSet, totalRefundedSet','Totales, ticket y validaciones'],['Estados','displayFinancialStatus, displayFulfillmentStatus','Cobro y preparación'],['LineItem','id, title, quantity, variant, originalTotalSet, originalUnitPriceSet','Detalle, devolución y cambio'],['Customer','id, firstName, lastName, email, phone','Identificación y ticket'],['Dirección','address1, city, province, country, zip','Lectura en detalle'],['Transaction','id, kind, status, gateway, amountSet','Pago manual y confirmación'],['FulfillmentOrder','id, status, supportedActions, remainingQuantity','Preparación tras cobro']], [1.25,3.45,2.15])
para('Métodos manejados: CASH, CARD, BIZUM, VOUCHER y MIXED. Shopify registra el pago como manual al marcar el pedido pagado; el desglose por método se conserva en el diario local, no como transacciones Shopify separadas.')

h('Reembolsos y cambios')
tbl(['Recurso u operación','Datos usados','Comportamiento'],[
['Refund','orderId, líneas lineItemId quantity restockType locationId, transacción parentId amount, note','Devuelve contra un pago manual exitoso; la nota incluye POS operación.'],['SuggestedRefund','importe, líneas, subtotal e impuesto','Calcula devolvible y evita exceder cantidades o saldo.'],['Location','id, name, isActive','Se elige si hay reposición.'],['Cambio','pedido original, devolución, nuevo DraftOrder y pedido sustituto','El crédito se aplica al nuevo pedido; diferencia se cobra o se emite en vale.'],['OrderCancel','orderId, motivo, refund false, restock','Solo después de reembolso completo desde la app.']], [1.55,3.25,2.08])

h('Clientes')
tbl(['Campo','Uso'],[['id','Relaciones con pedido, vale y devolución'],['firstName y lastName','Búsqueda y presentación'],['email','Búsqueda, asociación de vale y envío'],['phone','Búsqueda y alta'],['tags','En altas recuperables: pos-customer-identificador de operación']], [2.0,4.9])
para('La búsqueda devuelve hasta 10 clientes. Para emitir un vale de devolución o enviar un vale, la app asocia un cliente existente o crea uno. En altas de devolución exige nombre y correo o teléfono con prefijo internacional.')

h('Vales')
tbl(['Recurso','Campos Shopify usados','Regla de negocio'],[
['GiftCard','id, lastCharacters, maskedCode, initialValue, balance, enabled, expiresOn, createdAt','Saldo determina ACTIVE, EXHAUSTED o CANCELLED; código completo solo al crear.'],['Cliente','customer firstName lastName email','Asociación cuando existe.'],['Pedido','order name','Pedido de origen cuando Shopify lo relaciona.'],['note','Texto y bloque JSON POS_META','Nota visible y metadatos customerName, customerEmail, operationId.'],['GiftCardTransaction','id, amount, currencyCode, note, processedAt','Créditos y débitos; importe negativo equivale a canje.'],['giftCardDebit','debitAmount EUR y note','Descuenta el vale al cobrar.'],['giftCardDeactivate','id','Cancela sin eliminar.']], [1.35,3.25,2.28])
para('Formato de metadatos en note: texto visible, una línea ---POS_META--- y un JSON. Este bloque permite recuperar datos cuando el vale no tiene cliente asociado. No debe editarse manualmente salvo que se mantenga JSON válido.')

h('Sesiones de caja')
para('Las sesiones se guardan como metaobjetos de aplicación con el tipo $app:pos_session. La app requiere exactamente una sesión con status OPEN antes de cobrar, devolver o cambiar artículos.')
tbl(['Clave','Formato','Cuándo se escribe'],[
['cashier_name','Texto','Apertura; por defecto Cajero'],['opening_amount','Número serializado','Apertura'],['closing_amount','Número serializado','Apertura como 0 y cierre'],['expected_amount','Número serializado','Apertura como 0 y cierre calculado'],['difference','Número serializado','Apertura como 0 y cierre calculado'],['status','OPEN o CLOSED','Apertura, cierre o forzado'],['closed_at','ISO 8601 o vacío','Cierre'],['notes','Texto','Apertura, cierre o forzado'],['opened_at','Opcional, no se escribe actualmente','Si falta usa fecha local recordada o updatedAt mientras está abierta']], [2.15,1.65,3.05])

h('Diario local y límites de responsabilidad')
para('POS_DATA_DIR contiene journal.json en un volumen persistente. No es una base de datos de Shopify y no se sincroniza automáticamente con ella. Es imprescindible para no repetir mutaciones cuando se pierde una respuesta o se reinicia el servidor.')
tbl(['Colección local','Contenido relevante','Por qué no está en Shopify'],[
['operations','Clave, input, pasos, resultados, GIDs de sesión y vale','Pasos recuperables e idempotentes.'],['movements','orderId, sesión, tipo, método, importe, efectivo, cambio, pagos mixtos y cambio','Desglose de caja del TPV.'],['sessionDates','Fecha de apertura por GID de sesión','No confundir updatedAt de cierre con apertura.'],['mail jobs y configuración','Cola y configuración de correo propio','SMTP del TPV, no recurso Shopify.']], [1.55,3.55,1.75])
para('Consecuencia operativa: restaurar o borrar este volumen puede afectar a la conciliación de caja y a la recuperación de operaciones, aunque pedidos, vales y reembolsos sigan existiendo en Shopify.')

h('Configuración y permisos Shopify')
para('La integración usa Shopify Admin GraphQL API 2025 04. Accede con token de servidor; las credenciales no llegan al navegador. Los scopes configurados cubren productos, pedidos, borradores, clientes, inventario, vales, ubicaciones, cumplimientos y metaobjetos.')
tbl(['Área','Scopes configurados'],[
['Catálogo e inventario','read_products, write_products, read_inventory, write_inventory'],['Pedidos','read_orders, write_orders, read_draft_orders, write_draft_orders'],['Clientes','read_customers, write_customers'],['Vales','read_gift_cards, write_gift_cards, write_gift_card_transactions'],['Ubicaciones y fulfillment','read_locations, read_fulfillments, write_fulfillments y permisos de fulfillment'],['Metaobjetos','read_metaobject_definitions, write_metaobject_definitions, read_metaobjects, write_metaobjects']], [2.25,4.65])

h('Puntos de mantenimiento')
for x in ['No cambies el tipo categoria ni las claves nombre, imagen y padre sin actualizar el árbol de categorías.','No cambies el tipo $app:pos_session ni sus claves: la caja depende de esos nombres exactos.','Mantén Product custom.categorias como un array JSON de IDs de categoría.','Evita borrar las tags POS MML y pos-op de pedidos o borradores en curso.','Conserva el volumen POS_DATA_DIR antes de actualizar o migrar el servidor.']: bullet(x)
h('Fuente técnica')
para('Documento generado a partir del código actual del repositorio POSMML: consultas GraphQL de productos, pedidos y clientes; rutas de sesiones, vales, pagos, devoluciones y dashboard; y servicios de POS, contabilidad, cambios y recuperación. Fecha de revisión: 24 de septiembre de 2026.')
for s in doc.sections:
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('POS MML  •  Mapa de datos Shopify  •  Uso interno'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
doc.save(out); print(out)
